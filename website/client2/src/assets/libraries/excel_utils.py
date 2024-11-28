import datetime
import os
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook
from openpyxl.chart import (
    LineChart,
    BarChart,
    Reference,
)
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.layout import Layout, ManualLayout


class ExcelUtils:
    @staticmethod
    def create_line_chart(data_worksheet, series_info, row_end, start_column, end_column, chart_location, y_title):
        # Chart with date axis
        chart = LineChart()
        chart.title = y_title
        chart.style = 10
        chart.legend.position = 'b'

        chart.layout = Layout(
            ManualLayout(
                x=0.25,
                y=0.25,
                h=0.5,
                w=0.8,
                xMode="edge",
                yMode="edge",
            ))

        # vertical
        chart.y_axis.crossAx = 500
        y_values = Reference(
            data_worksheet,
            min_col=start_column,
            max_col=end_column,
            min_row=1,
            max_row=row_end)

        # horizontal
        chart.x_axis = DateAxis(crossAx=100)
        x_values = Reference(
            data_worksheet,
            min_col=series_info['series_column'],
            max_col=series_info['series_column'],
            min_row=2,
            max_row=row_end)

        # build the chart
        chart.add_data(y_values, titles_from_data=True)
        chart.set_categories(x_values)
        data_worksheet.add_chart(chart, chart_location)

    @staticmethod
    def excel_find_last_cell(data_worksheet, column_letter):
        # find the first available cell
        last_cell = 0
        for cell in data_worksheet[column_letter]:
            if cell.value is None:
                last_cell = cell.row
                break
            else:
                last_cell = cell.row + 1

        return last_cell
